"""Opt-in local AI and n8n-container SMTP contract proof using synthetic demo data."""

import argparse
import io
import json
import subprocess
import sys
from pathlib import Path
from uuid import uuid4

import httpx
from docx import Document
from minori.core.config import PROJECT_ROOT, get_settings
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from verify_ai_gateway import SYNTHETIC_SOURCE


def documents():
    doc = Document()
    doc.add_paragraph(SYNTHETIC_SOURCE)
    buffer = io.BytesIO()
    doc.save(buffer)
    yield "vacancy.docx", buffer.getvalue()
    workbook = Workbook()
    workbook.active.append(["求人", "企業"])
    workbook.active.append([SYNTHETIC_SOURCE, "架空企業"])
    workbook.active.append(["第二の架空求人。JFT-Basic A2.2以上。", "第二の架空企業"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    yield "vacancy.xlsx", buffer.getvalue()
    yield "vacancy.csv", ("求人\n" + SYNTHETIC_SOURCE + "\n第二の架空求人\n").encode()
    yield "vacancy.txt", SYNTHETIC_SOURCE.encode()
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 20 300 Td (Synthetic demo vacancy JLPT N4) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(contents)
    buffer = io.BytesIO()
    writer.write(buffer)
    yield "vacancy.pdf", buffer.getvalue()


NODE_WORKER = r"""
const net = require('node:net');
const crypto = require('node:crypto');
const config = CONFIG_PLACEHOLDER;
async function request(path, body) {
  const r = await fetch(config.api + path, {method: body ? 'POST' : 'GET',
    headers: {'Authorization': 'Bearer ' + config.token, 'Content-Type': 'application/json'},
    body: body ? JSON.stringify(body) : undefined, signal: AbortSignal.timeout(10000)});
  if (!r.ok) throw new Error('Backend HTTP ' + r.status);
  return await r.json();
}
async function smtp(message) {
  const socket = net.createConnection({host: config.mailpit, port: 1025});
  let buffer = '', waiting;
  const replies = [];
  socket.setTimeout(10000, () => socket.destroy(new Error('SMTP timeout')));
  socket.on('data', chunk => {
    buffer += chunk.toString();
    let match;
    while ((match = buffer.match(/(?:^|\r\n)(\d{3}) [^\r\n]*\r\n/))) {
      const end = match.index + match[0].length;
      replies.push(Number(match[1])); buffer = buffer.slice(end);
      if (waiting) { const done = waiting; waiting = undefined; done.resolve(replies.shift()); }
    }
  });
  let failed;
  socket.on('error', error => { failed = error; if (waiting) { waiting.reject(error); waiting = undefined; } });
  function reply() { if (failed) return Promise.reject(failed); if (replies.length) return Promise.resolve(replies.shift()); return new Promise((resolve, reject) => { waiting = {resolve, reject}; }); }
  async function command(value, expected) { socket.write(value + '\r\n'); if (await reply() !== expected) throw new Error('SMTP command rejected'); }
  try {
    if (await reply() !== 220) throw new Error('SMTP greeting rejected');
    await command('EHLO minori.local', 250);
    await command('MAIL FROM:<pengingat@minori.local>', 250);
    await command('RCPT TO:<mailpit@minori.local>', 250);
    await command('DATA', 354);
    socket.write(message + '\r\n.\r\n');
    if (await reply() !== 250) throw new Error('SMTP DATA rejected');
    await command('QUIT', 221);
  } finally { socket.destroy(); }
}
(async () => {
  const rows = await request('/internal/reminders/due?limit=10');
  const processed = [];
  for (const row of rows) {
    const claim = await request('/internal/reminders/claim', {notification_id: row.id});
    if (!claim) continue;
    const notification = claim.notification, routing = claim.recipient_routing;
    if (notification.dry_run) {
      if (routing.delivery_allowed || routing.effective_recipient) throw new Error('Unsafe simulation routing');
      await request('/internal/reminders/' + notification.id + '/complete', {claim_token: claim.claim_token, outcome: 'simulated'});
      processed.push({template: notification.template_code, outcome: 'simulated'}); continue;
    }
    if (!routing.delivery_allowed || routing.routing_mode !== 'mailpit' || routing.effective_recipient !== 'mailpit@minori.local') throw new Error('SMTP routing safety validation failed');
    const messageId = crypto.randomUUID() + '@minori.local';
    const message = ['From: MINORI Placement Intelligence <pengingat@minori.local>',
      'To: mailpit@minori.local', 'Subject: =?UTF-8?B?' + Buffer.from(claim.message.subject).toString('base64') + '?=',
      'Message-ID: <' + messageId + '>', 'MIME-Version: 1.0', 'Content-Type: text/plain; charset=UTF-8',
      'Content-Transfer-Encoding: base64', '', Buffer.from(claim.message.body).toString('base64')].join('\r\n');
    // Fixed local SMTP only. Once SMTP succeeds, never report a false delivery failure.
    await smtp(message);
    const r = await fetch('http://' + config.mailpit + ':8025/api/v1/messages');
    const mailbox = await r.json();
    const captured = mailbox.messages.find(x => x.MessageID === messageId);
    if (!captured || captured.To.length !== 1 || captured.To[0].Address !== routing.effective_recipient) throw new Error('Mailpit receipt not verified; callback must be reviewed');
    const callback = {claim_token: claim.claim_token, outcome: 'sent', provider_reference: 'mailpit:' + captured.ID};
    const completed = await request('/internal/reminders/' + notification.id + '/complete', callback);
    await request('/internal/reminders/' + notification.id + '/complete', callback);
    if (completed.status !== 'sent') throw new Error('Callback did not record SMTP success');
    processed.push({template: notification.template_code, outcome: 'sent', mailpit_id: captured.ID,
      notification_id: notification.id, effective_recipient: routing.effective_recipient,
      approved_snapshot: Boolean(notification.approved_message)});
  }
  console.log(JSON.stringify({passed: true, scope: 'contract probe inside n8n container; existing workflow unchanged', initial_due: rows.length, processed}));
})().catch(() => { console.error('Contract probe failed. Inspect queue and Mailpit before retrying SMTP.'); process.exitCode = 1; });
"""


def verify(mode):
    settings = get_settings()
    if mode == "test":
        from dotenv import dotenv_values

        token = dotenv_values(PROJECT_ROOT / ".env")["AI_TEST_API_TOKEN"]
        url, internal_url, mailpit = (
            "http://127.0.0.1:18100",
            "http://backend-ai-test:8000",
            "mailpit-ai-test",
        )
    else:
        token = settings.internal_api_token.get_secret_value()
        url, internal_url, mailpit = "http://127.0.0.1:8000", "http://backend:8000", "mailpit-demo"
    with httpx.Client(
        base_url=url, headers={"Authorization": "Bearer " + token}, timeout=300
    ) as client:

        def call(method, path, **kwargs):
            response = client.request(method, path, **kwargs)
            if not response.is_success:
                raise RuntimeError(f"HTTP {response.status_code} for {path}; body omitted")
            return response.json()

        candidates = call("GET", "/candidates")
        assert len(candidates) == 24 and all(c["email"].endswith(".invalid") for c in candidates)
        formats = []
        for filename, data in documents():
            preview = call("POST", "/ai/documents/preview", files={"file": (filename, data)})
            assert preview["rows"]
            formats.append(preview["kind"])
        analysis = call("POST", "/ai/vacancies/analyze", json={"text": SYNTHETIC_SOURCE})
        assert (
            analysis["extraction"]["quota"] == 3
            and analysis["translation"]
            and analysis["review"]["mandatory"]
        )
        vacancy = next(v for v in call("GET", "/vacancies") if v["reference"] == "DEMO-PRO-VAC-01")
        candidate = next(c for c in candidates if c["reference"] == "DEMO-PRO-CAND-01")
        path = f"/matching/candidates/{candidate['id']}/vacancies/{vacancy['id']}/explanation"
        official = call("GET", path)
        explained = call(
            "POST",
            "/ai/matching/explanation",
            json={"candidate_id": candidate["id"], "vacancy_id": vacancy["id"]},
        )
        assert explained["official"]["total_score"] == official["total_score"]
        unchanged = call("GET", path)
        for key in (
            "eligibility_status",
            "total_score",
            "input_fingerprint",
            "explanation",
            "staff_status",
        ):
            assert unchanged[key] == official[key]
        rows = call("GET", "/reminders")
        due = call("GET", "/internal/reminders/due")
        future = next(
            row
            for row in rows
            if row["status"] == "pending" and row["id"] not in {x["id"] for x in due}
        )
        preview = call("POST", f"/reminders/{future['id']}/email/draft")
        assert preview["source"] == "ai" and preview["facts"]["scheduled_wib"] in preview["body"]
        payload = {
            "subject": preview["subject"] + " (Ditinjau staf)",
            "body": preview["body"] + "\n\nPesan telah ditinjau staf.",
            "facts_fingerprint": preview["facts_fingerprint"],
            "idempotency_key": uuid4().hex,
            "confirmed": True,
        }
        approved = call("POST", f"/reminders/{future['id']}/email/approve", json=payload)
        assert approved["approved_message"]["body"] == payload["body"]
        assert (
            call("POST", f"/reminders/{future['id']}/email/approve", json=payload)["id"]
            == future["id"]
        )
        script = NODE_WORKER.replace(
            "CONFIG_PLACEHOLDER",
            json.dumps({"api": internal_url, "mailpit": mailpit, "token": token}),
        )
        completed = subprocess.run(
            ["docker", "exec", "-i", "minori-placement-intelligence-n8n-1", "node", "-"],
            input=script,
            capture_output=True,
            text=True,
            timeout=120,
        )
        if completed.returncode:
            raise RuntimeError(
                "n8n-container contract probe failed; review claim/Mailpit before retrying"
            )
        delivery = json.loads(completed.stdout)
        assert any(x.get("approved_snapshot") for x in delivery["processed"])
        assert call("GET", f"/reminders/{future['id']}")["status"] == "sent"
        return {
            "passed": True,
            "model": settings.nine_router_model,
            "document_formats": formats,
            "ai_extraction_translation_review": True,
            "ai_matching_official_score_preserved": True,
            "ai_email_approved_and_sent": True,
            "delivery": delivery,
        }


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=["test", "demo"], default="test")
    parser.add_argument("--output", type=Path, default=PROJECT_ROOT / "test-results/ai-e2e.json")
    args = parser.parse_args()
    result = verify(args.mode)
    args.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
