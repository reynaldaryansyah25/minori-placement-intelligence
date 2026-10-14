"""Read bounded text documents in memory; never extract archives or execute formulas."""

import csv
import io
import zipfile
from pathlib import PurePosixPath

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT = 12000
MAX_ROWS = 100


class DocumentError(ValueError):
    pass


def _decode(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp932"):
        try:
            text = data.decode(encoding)
            if "\x00" in text:
                raise DocumentError("Dokumen mengandung data biner.")
            return text
        except UnicodeDecodeError:
            continue
    raise DocumentError("Encoding teks tidak didukung. Gunakan UTF-8.")


def _zip_safe(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(x.file_size for x in entries) > 20 * 1024 * 1024:
            raise DocumentError("Arsip dokumen melebihi batas aman.")
        for item in entries:
            path = PurePosixPath(item.filename)
            if path.is_absolute() or ".." in path.parts or item.flag_bits & 1:
                raise DocumentError("Arsip dokumen tidak aman.")
            if any(x in item.filename.lower() for x in ("vbaproject", "/embeddings/", "/activex/")):
                raise DocumentError("Dokumen dengan makro atau objek tertanam tidak didukung.")
            if item.filename.endswith(".rels") and b'TargetMode="External"' in archive.read(item):
                raise DocumentError("Dokumen dengan referensi eksternal tidak didukung.")
            if item.filename.endswith((".xml", ".rels")):
                xml = archive.read(item).upper()
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise DocumentError("Dokumen dengan deklarasi entitas tidak didukung.")


def _rows(values, kind):
    if not values:
        raise DocumentError("Dokumen tidak memuat teks. PDF scan memerlukan OCR terlebih dahulu.")
    if len(values) > MAX_ROWS:
        raise DocumentError("Maksimal 100 baris lowongan per dokumen.")
    rows = []
    for number, text in enumerate(values, 1):
        text = text.strip()
        if not text or len(text) > MAX_TEXT:
            raise DocumentError(
                "Setiap lowongan harus berisi 1–12.000 karakter. Pisahkan dokumen panjang."
            )
        rows.append({"number": number, "text": text, "preview": text[:240]})
    return {"rows": rows, "kind": kind, "warning": None}


def parse_document(filename, data):
    if not data or len(data) > MAX_FILE_BYTES:
        raise DocumentError("Ukuran dokumen harus antara 1 byte dan 5 MB.")
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if suffix not in {".txt", ".csv", ".pdf", ".docx", ".xlsx"}:
        raise DocumentError("Gunakan PDF, DOCX, XLSX, CSV, atau TXT.")
    try:
        if suffix in {".txt", ".csv"}:
            text = _decode(data)
            if suffix == ".txt":
                return _rows([text], "txt")
            reader = (
                csv.reader(io.StringIO(text), delimiter=";")
                if text.count(";") > text.count(",")
                else csv.reader(io.StringIO(text))
            )
            matrix = []
            for row in reader:
                if len(row) > 30 or any(len(x) > 4000 for x in row):
                    raise DocumentError("Maksimal 30 kolom dan 4.000 karakter per sel.")
                if any(x.strip() for x in row):
                    matrix.append(row)
                if len(matrix) > MAX_ROWS + 1:
                    raise DocumentError("Maksimal 100 baris lowongan.")
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise DocumentError("Isi dokumen tidak sesuai format PDF.")
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted or len(reader.pages) > 30:
                raise DocumentError("PDF terenkripsi atau lebih dari 30 halaman tidak didukung.")
            text = ""
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
                if len(text) > MAX_TEXT:
                    raise DocumentError("Teks PDF terlalu panjang. Pilih bagian yang diperlukan.")
            return _rows([text] if text.strip() else [], "pdf")
        else:
            _zip_safe(data)
            if suffix == ".docx":
                doc = Document(io.BytesIO(data))
                parts = [x.text for x in doc.paragraphs]
                parts.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for table in doc.tables
                    for row in table.rows
                )
                return _rows(["\n".join(parts)], "docx")
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
            try:
                sheet = workbook.active
                if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 30:
                    raise DocumentError("Maksimal 100 baris lowongan dan 30 kolom per lembar.")
                matrix = []
                for row in sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=31, values_only=True):
                    if row[-1] is not None:
                        raise DocumentError("Maksimal 30 kolom per lembar.")
                    cells = [str(x) if x is not None else "" for x in row[:-1]]
                    if any(len(x) > 4000 for x in cells):
                        raise DocumentError("Sel Excel terlalu panjang.")
                    if any(cells):
                        matrix.append(cells)
                if len(workbook.worksheets) != 1:
                    raise DocumentError("Gunakan satu lembar lowongan per file Excel.")
            finally:
                workbook.close()
        if len(matrix) < 2:
            raise DocumentError("Tabel memerlukan judul kolom dan minimal satu baris data.")
        headers = matrix[0]
        texts = [
            "\n".join(
                f"{header}: {value}"
                for header, value in zip(headers, row, strict=False)
                if value.strip()
            )
            for row in matrix[1:]
        ]
        return _rows(texts, suffix[1:])
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError(
            "Dokumen rusak atau formatnya tidak sesuai. Ekspor ulang dokumen."
        ) from error
